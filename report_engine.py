"""Сопоставление дневных выгрузок с первым листом КЕТИК."""
import io
import re
import posixpath
import zipfile
import xml.etree.ElementTree as ET
from collections import defaultdict
from copy import copy
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

def clean(value):
    return re.sub(r'\s+', ' ', str(value if value is not None else '')).strip()

def key(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return clean(value).casefold()

def number(value):
    s = re.sub(r'\s+', '', str(value)).replace(',', '.').replace('−', '-')
    try:
        result = Decimal(s)
        if not result.is_finite():
            raise InvalidOperation
        return result
    except InvalidOperation:
        raise ValueError(f'Некорректное или пустое число: {value!r}')

def raw_rows(data):
    # Выгрузка программы содержит некорректные стили. Читаем значения XML.
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        strings = []
        if 'xl/sharedStrings.xml' in z.namelist():
            strings = [''.join(t.text or '' for t in si.findall('.//m:t', NS))
                       for si in ET.fromstring(z.read('xl/sharedStrings.xml'))]
        book = ET.fromstring(z.read('xl/workbook.xml'))
        first = book.find('m:sheets/m:sheet', NS)
        rid = first.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        target = next(r.get('Target') for r in rels if r.get('Id') == rid)
        path = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
        result = []
        for row in ET.fromstring(z.read(path)).findall('m:sheetData/m:row', NS):
            values = {}
            for c in row.findall('m:c', NS):
                col = openpyxl.utils.cell.column_index_from_string(re.match('[A-Z]+', c.get('r')).group())
                v = c.find('m:v', NS)
                value = v.text if v is not None else ''
                if c.get('t') == 's':
                    value = strings[int(value)]
                elif c.get('t') == 'inlineStr':
                    value = ''.join(t.text or '' for t in c.findall('.//m:t', NS))
                values[col] = value
            result.append((int(row.get('r')), values))
        return result

def parse_sales(data, name):
    rows = raw_rows(data)
    header_pos = next((i for i, (_, r) in enumerate(rows)
                       if 'артикул' in map(key, r.values()) and 'доход' in map(key, r.values())), None)
    if header_pos is None:
        raise ValueError(f'{name}: не найдены заголовки Артикул, Кол., Доход, Прибыль.')
    dates = set()
    for _, row in rows[:header_pos]:
        for value in row.values():
            dates.update(re.findall(r'\b\d{2}\.\d{2}\.\d{4}\b', str(value)))
    if not dates:
        dates.update(re.findall(r'\b\d{2}\.\d{2}\.\d{4}\b', name))
    if len(dates) != 1:
        raise ValueError(f'{name}: нужна выгрузка за один день с полной датой. Период нельзя распределить по дням.')
    day = datetime.strptime(dates.pop(), '%d.%m.%Y').date()
    columns = {}
    for c, v in rows[header_pos][1].items():
        columns.setdefault(key(v), c)  # Повторенные подписи объединенных ячеек.
    qty = next((columns[k] for k in ('кол.', 'количество', 'кол') if k in columns), None)
    if qty is None or any(k not in columns for k in ('артикул', 'доход', 'прибыль')):
        raise ValueError(f'{name}: отсутствуют обязательные столбцы.')
    metrics = [qty, columns['доход'], columns['прибыль']]
    items = defaultdict(lambda: [Decimal(0), Decimal(0), Decimal(0)])
    control = None
    for rownum, row in rows[header_pos + 1:]:
        if any(key(v).startswith('итого') for c, v in row.items() if c < qty):
            if all(clean(row.get(c)) for c in metrics):
                control = [number(row[c]) for c in metrics]
            break
        art = key(row.get(columns['артикул']))
        if not art or art == 'артикул':
            continue
        try:
            vals = [number(row.get(c, '')) for c in metrics]
        except ValueError as e:
            raise ValueError(f'{name}, строка {rownum}: {e}') from e
        items[art] = [a + b for a, b in zip(items[art], vals)]
    if not items:
        raise ValueError(f'{name}: нет товарных строк.')
    total = [sum(v[i] for v in items.values()) for i in range(3)]
    if control and any(abs(a-b) > Decimal('0.01') for a, b in zip(total, control)):
        raise ValueError(f'{name}: суммы товаров не совпадают со строкой Итого: {total} / {control}.')
    return day, dict(items)

def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(clean(value), '%d.%m.%Y').date()
    except ValueError:
        return None

def ensure_column(ws, label):
    for cell in ws[1]:
        if key(cell.value) == key(label):
            return cell.column
    col = ws.max_column + 1
    ws.cell(1, col, label)
    ws.cell(1, col)._style = copy(ws.cell(1, max(1, col-1))._style)
    ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = max(18, len(label)+2)
    return col

def build_report(base, files):
    wb = openpyxl.load_workbook(io.BytesIO(base))
    ws = wb.worksheets[0]
    headers = {key(c.value): c.column for c in ws[1] if c.value is not None}
    if not {'категория', 'артикул'} <= headers.keys():
        raise ValueError('На первом листе в первой строке нужны Категория и Артикул.')
    articles, categories = {}, {}
    for r in range(2, ws.max_row + 1):
        art = key(ws.cell(r, headers['артикул']).value)
        if not art:
            continue
        if art in articles:
            raise ValueError(f'Артикул {art.upper()} повторяется на первом листе. Удалите неоднозначность перед расчетом.')
        category = clean(ws.cell(r, headers['категория']).value) or 'Без категории'
        categories.setdefault(key(category), category)
        articles[art] = (r, key(category))
    daily = {}
    for name, data in files:
        day, items = parse_sales(data, name)
        if day in daily:
            raise ValueError(f'Загружены две выгрузки за {day:%d.%m.%Y}. Оставьте одну полную выгрузку за день, чтобы исключить двойной учет.')
        daily[day] = items
    if not daily:
        raise ValueError('Загрузите хотя бы одну выгрузку.')
    unknown = sorted({a for items in daily.values() for a in items if a not in articles})
    if unknown:
        raise ValueError('На первом листе не найдены артикулы: ' + ', '.join(a.upper() for a in unknown) + '. Добавьте их с категориями и загрузите файл повторно.')
    summary = {cat: [Decimal(0)]*3 for cat in categories}
    article_totals = {a: [Decimal(0)]*3 for a in articles}
    report = wb['ОТЧЕТ'] if 'ОТЧЕТ' in wb.sheetnames else wb.create_sheet('ОТЧЕТ')
    if report.cell(1,1).value is None:
        report.cell(1,1,'ДАТА')
    cat_cols = {cat: ensure_column(report, label) for cat, label in categories.items()}
    count_col = ensure_column(report, 'ОБЩЕЕ КОЛИЧЕСТВО')
    income_col = ensure_column(report, 'ОБЩИЙ Доход')
    profit_col = ensure_column(report, 'прибыль')
    for day, items in sorted(daily.items()):
        date_col = next((c.column for c in ws[1] if as_date(c.value) == day), None)
        if date_col is None:
            date_col = ensure_column(ws, day.strftime('%d.%m.%Y'))
        row = next((r for r in range(2, report.max_row+1) if as_date(report.cell(r,1).value) == day), None)
        if row is None:
            row = report.max_row+1
            report.cell(row,1,day).number_format = 'dd.mm.yyyy'
        for col in cat_cols.values():
            report.cell(row,col,0)
        totals = [Decimal(0)]*3
        for art, (r, cat) in articles.items():
            values = items.get(art, [Decimal(0)]*3)
            ws.cell(r,date_col,float(values[0]))
            article_totals[art] = [a+b for a,b in zip(article_totals[art], values)]
            summary[cat] = [a+b for a,b in zip(summary[cat], values)]
            cell = report.cell(row,cat_cols[cat])
            cell.value += float(values[0])
            totals = [a+b for a,b in zip(totals, values)]
        for col, val in zip((count_col,income_col,profit_col),totals):
            report.cell(row,col,float(val)).number_format = '#,##0.00'
    # Существующие Доход/Прибыль не имеют подтвержденного периода.
    # Отдельные явно подписанные поля содержат только загруженный период.
    out_cols = [ensure_column(ws, label) for label in
                ('Количество за загруженные дни', 'Доход за загруженные дни', 'Прибыль за загруженные дни')]
    date_cols = [c.column for c in ws[1] if as_date(c.value)]
    for art, (r, _) in articles.items():
        for col, val in zip(out_cols,article_totals[art]):
            ws.cell(r,col,float(val)).number_format = '#,##0.00'
        if 'общ.кол' in headers:
            refs = ','.join(f'{openpyxl.utils.get_column_letter(c)}{r}' for c in date_cols)
            ws.cell(r,headers['общ.кол'],f'=SUM({refs})')
    title = 'Итоги загруженных дней'
    if title in wb.sheetnames:
        del wb[title]
    out = wb.create_sheet(title)
    out.append(['Период', ', '.join(d.strftime('%d.%m.%Y') for d in sorted(daily))])
    out.append(['Категория','Общее количество','Общий доход','Общая прибыль'])
    records = []
    for cat, vals in summary.items():
        record = [categories[cat]] + [float(v) for v in vals]
        records.append(record)
        out.append(record)
    total_row = out.max_row+1
    out.cell(total_row,1,'ИТОГО')
    for col in range(2,5):
        letter = openpyxl.utils.get_column_letter(col)
        out.cell(total_row,col,f'=SUM({letter}3:{letter}{total_row-1})')
    for row in out.iter_rows(min_row=2):
        for cell in row:
            if cell.row in (2,total_row):
                cell.font = Font(bold=True,color='FFFFFF')
                cell.fill = PatternFill('solid',fgColor='1E3A8A')
            if cell.column > 1 and cell.row > 2:
                cell.number_format = '#,##0.00'
    for col,width in [('A',42),('B',26),('C',24),('D',24)]:
        out.column_dimensions[col].width = width
    out.cell(1,2).alignment = Alignment(wrap_text=True)
    out.row_dimensions[1].height = max(30, 15*((len(files)+1)//2))
    out.freeze_panes = 'B3'
    out.auto_filter.ref = f'A2:D{total_row-1}'
    wb.active = wb.sheetnames.index(title)
    output = io.BytesIO()
    wb.save(output)
    totals = [float(sum(v[i] for v in summary.values())) for i in range(3)]
    return output.getvalue(), records, totals
